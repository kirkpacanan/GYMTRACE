#!/usr/bin/env python3
"""Build GYMTRACE M1|S2 Software Project Proposal (IEEE A4, 2-column)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "se_manuscript_assets"
OUT = Path.home() / "Downloads" / "GYMTRACE_Software_Project_Proposal_Manuscript.docx"


def set_run_font(run, size=10, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_two_columns(section, num=2, space_twips="360"):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), space_twips)


def apply_ieee_page(section):
    # Geometry mirrored from conference-template-a4
    section.page_width = Pt(595.30)
    section.page_height = Pt(841.90)
    section.top_margin = Pt(27)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(44.65)
    section.right_margin = Pt(44.65)
    section.header_distance = Pt(36)
    section.footer_distance = Pt(36)


def add_para(doc, text, *, size=10, bold=False, italic=False, align="left", space_after=6, space_before=0, first_line=True):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line and align == "justify":
        pf.first_line_indent = Pt(18)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_mixed_para(doc, parts, *, align="justify", space_after=6, first_line=True):
    """parts: list of (text, bold, italic, size)"""
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line and align == "justify":
        pf.first_line_indent = Pt(18)
    for text, bold, italic, size in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_ieee(doc, text, level=1):
    if level == 1:
        return add_para(doc, text, size=10, bold=True, align="center", space_before=10, space_after=8, first_line=False)
    if level == 2:
        return add_para(doc, text, size=10, bold=True, italic=True, align="left", space_before=8, space_after=4, first_line=False)
    return add_para(doc, text, size=10, bold=True, align="left", space_before=6, space_after=3, first_line=False)


def add_figure(doc, path: Path, caption: str, width=Inches(3.35)):
    if not path.exists():
        add_para(doc, f"[Missing figure: {path.name}]", italic=True, first_line=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=8, bold=True)


def add_usecase_table(doc, rows):
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Description"
    for cell in hdr:
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=8, bold=True)
    for i, (field, desc) in enumerate(rows):
        table.rows[i + 1].cells[0].text = field
        table.rows[i + 1].cells[1].text = desc
        for cell in table.rows[i + 1].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=8)
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    apply_ieee_page(section)
    set_two_columns(section, num=1)  # title block single column first

    # Normal style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # ---- Title block (single column) ----
    add_para(
        doc,
        "GYMTRACE: A Facial Recognition–Based Software System for Automated Gym Membership Authentication and Occupancy Prediction",
        size=18,
        bold=True,
        align="center",
        space_after=10,
        first_line=False,
    )

    authors = [
        ("Nicholas Klein Castillanes", "nkCastillanes@mcm.edu.ph"),
        ("Kirk Roden Pacanan", "krPacanan@mcm.edu.ph"),
        ("Gary Louise Querequincia", "glQuerequincia@mcm.edu.ph"),
    ]
    for name, email in authors:
        add_para(doc, name, size=11, bold=True, align="center", space_after=0, first_line=False)
        add_para(doc, "College of Computer and Information Science", size=9, align="center", space_after=0, first_line=False)
        add_para(doc, "Mapua Malayan Colleges Mindanao", size=9, align="center", space_after=0, first_line=False)
        add_para(doc, "Davao City, Philippines", size=9, align="center", space_after=0, first_line=False)
        add_para(doc, email, size=9, italic=True, align="center", space_after=8, first_line=False)

    # Switch to 2-column before Abstract (IEEE: title/authors full-width; abstract+body two-column)
    body = doc.add_section(WD_SECTION.CONTINUOUS)
    apply_ieee_page(body)
    set_two_columns(body, num=2, space_twips="360")

    add_mixed_para(
        doc,
        [
            ("Abstract—", True, False, 9),
            (
                "Gym facilities still rely on membership cards, QR codes, or manual desk checks that are easy to share, slow during peak hours, and weak at producing trustworthy occupancy data. "
                "This software project proposal presents GYMTRACE, a web-based campus gym system that authenticates members through facial recognition, supports day-pass visitors, and uses verified attendance events to drive short-horizon occupancy prediction for members, staff, and administrators. "
                "The proposal follows an object-oriented software engineering approach with UML use-case, activity, class, and sequence models plus a three-tier system architecture. "
                "A working Next.js prototype implements role-based interfaces for member check-in, day-pass purchase/entry, staff enrollment and logs, and an admin occupancy dashboard. "
                "The contribution is an integrated software blueprint that couples contactless access control with occupancy analytics for campus and commercial gym operations.",
                False,
                False,
                9,
            ),
        ],
        align="justify",
        first_line=False,
        space_after=6,
    )
    add_mixed_para(
        doc,
        [
            ("Keywords—", True, False, 9),
            (
                "facial recognition, gym membership authentication, occupancy prediction, UML, three-tier architecture, biometric access control, software project proposal",
                False,
                False,
                9,
            ),
        ],
        align="justify",
        first_line=False,
        space_after=10,
    )

    # ===================== SECTION I =====================
    add_heading_ieee(doc, "I. INTRODUCTION")

    add_heading_ieee(doc, "A. Nature of the System and Company", level=2)
    add_para(
        doc,
        "GYMTRACE is a proposed campus gym operations software product developed as an academic software engineering and machine-learning project by students of the College of Computer and Information Science at Mapua Malayan Colleges Mindanao. "
        "The “company” context for this proposal is a campus or commercial fitness facility that needs secure member entry, temporary day-pass handling, and data-informed capacity management. "
        "The system is delivered as a progressive web application with kiosk-style check-in screens, staff consoles, and an administrator dashboard. "
        "Core capabilities are (1) contactless member authentication via facial recognition, (2) day-pass purchase and validation, (3) attendance logging, and (4) short-horizon occupancy forecasting displayed to members and managers.",
        align="justify",
    )

    add_heading_ieee(doc, "B. Problem Statement", level=2)
    add_para(
        doc,
        "Traditional gym authentication methods—membership cards, mobile QR codes, and staff-assisted ID checks—are convenient in principle but fragile in practice. "
        "Credentials can be shared, lost, or forged; queues form at entry points during rush periods; and front-desk verification does not reliably produce continuous occupancy records that managers can use for staffing and capacity control. "
        "At the same time, members lack a trustworthy view of how crowded the gym will be in the next half hour, which reduces service quality and can create unsafe peak loads. "
        "GYMTRACE therefore addresses an integrated software problem: design and implement a secure, role-based gym entry system whose authenticated events also feed occupancy prediction services for operational decision support.",
        align="justify",
    )

    add_heading_ieee(doc, "C. Objectives", level=2)
    add_para(doc, "The general objective is to design and propose GYMTRACE as a software system for automated gym membership authentication and occupancy prediction.", align="justify")
    add_para(doc, "Specific objectives are to:", align="justify", first_line=False)
    objectives = [
        "Analyze stakeholder needs for members, day-pass users, staff, and administrators and translate them into UML requirements models.",
        "Design use-case, activity, class, sequence, and three-tier architecture views that specify how facial authentication, attendance logging, and occupancy services interact.",
        "Implement a prototype graphical user interface that demonstrates check-in, day-pass, staff, and admin workflows.",
        "Integrate an occupancy prediction pipeline (baseline Linear Regression and Random Forest) so authenticated presence can support short-horizon forecasts.",
        "Document scope, limitations, and related work to guide further development and evaluation.",
    ]
    for i, obj in enumerate(objectives, 1):
        add_para(doc, f"({i}) {obj}", align="justify", first_line=False, space_after=3)

    add_heading_ieee(doc, "D. Scope and Limitation", level=2)
    add_para(
        doc,
        "Scope. GYMTRACE covers enrolled members and staff within a single gym site; camera-based verification at designated entry points; day-pass purchase and entry validation; attendance logging; member occupancy viewing; staff member/day-pass management; and admin occupancy dashboards and reports. "
        "The software architecture is a three-tier web system (presentation, application services, and data stores) with machine-learning components for face matching and occupancy forecasting.",
        align="justify",
    )
    add_para(
        doc,
        "Limitations. Multi-branch enterprise analytics, full payment-gateway settlement, building HVAC optimization, and production-grade biometric hardware certification are outside the present proposal. "
        "The current prototype uses a mock face-scan service suitable for demonstration; production deployment would require enrolled face galleries, privacy compliance, and calibrated matching thresholds. "
        "Occupancy models may be trained on public campus-gym historical data and later recalibrated on the facility’s own logs.",
        align="justify",
    )

    add_heading_ieee(doc, "E. Related Works", level=2)
    add_para(
        doc,
        "This subsection reviews ten recent works (2021–2026) on facial authentication systems, attendance/access software, and occupancy prediction platforms that inform GYMTRACE’s software design.",
        align="justify",
    )
    related = [
        (
            "Meng et al. proposed MagFace, a quality-aware deep face embedding method that improves recognition under uneven capture quality—relevant to gym cameras with pose and lighting variation [1].",
        ),
        (
            "Kim et al. introduced AdaFace, which adapts angular margins by image quality and strengthens open-set identification pipelines used in access-control software [2].",
        ),
        (
            "Boutros et al. presented ElasticFace, sampling elastic margins to improve separability under large intra-class variation common at uncontrolled entrances [3].",
        ),
        (
            "Chowanda et al. built a real-time employee attendance system with MTCNN detection, alignment, and FaceNet embeddings, emphasizing gallery maintenance—an applied software stack GYMTRACE adapts for membership verification [4].",
        ),
        (
            "Kushwaha et al. developed a CNN-based classroom attendance manager that automates roll call from live video and reduces proxy attendance, illustrating contactless logging UX patterns [5].",
        ),
        (
            "Ejaz et al. applied facial recognition to residence-hall entry control, matching arrivals against an enrolled gallery and alerting on unknown faces—closely aligned with gym door authentication [6].",
        ),
        (
            "Tekler and Chong evaluated deep occupancy predictors across office, library, and lecture spaces using minimum sensing, showing that sequence models can power facility dashboards [7].",
        ),
        (
            "Hitimana et al. implemented an IoT plus deep-learning occupancy framework, demonstrating how backend services can fuse sensor/event streams for building intelligence [8].",
        ),
        (
            "Kanthila et al. proposed cascaded LSTM occupancy prediction for buildings, motivating short-horizon forecast services behind admin and member occupancy screens [9].",
        ),
        (
            "Diarra et al. used LSTM models with non-intrusive cues to predict multi-room occupancy states, supporting the idea that authenticated gym events can feed predictive UIs [10].",
        ),
    ]
    for (text,) in related:
        add_para(doc, text, align="justify")
    add_para(
        doc,
        "Across these studies, authentication systems typically stop at identity confirmation, while occupancy systems often ignore verified membership events. "
        "GYMTRACE’s software proposal closes that gap by designing one product where face-authenticated entries both gate access and update occupancy services.",
        align="justify",
    )

    # ===================== SECTION II =====================
    add_heading_ieee(doc, "II. SOFTWARE DEVELOPMENT METHODOLOGY")
    add_para(
        doc,
        "GYMTRACE follows an iterative, object-oriented software engineering process: requirements modeling with UML, architectural design of a three-tier web system, and prototype implementation of role-based user interfaces. "
        "The diagrams below specify the main success and failure paths for gym entry and occupancy update.",
        align="justify",
    )

    # A. Use Case
    add_heading_ieee(doc, "A. UML Use Case Diagram with Use Case Description", level=2)
    add_para(
        doc,
        "Fig. 1 shows the GYMTRACE use-case model. Primary actors are Day-Pass User, Member, Staff, and Admin. "
        "Members authenticate via facial recognition (including membership verification and attendance logging), view attendance history, and view predicted occupancy. "
        "Day-pass users purchase a pass and enter the gym (including attendance logging). "
        "Staff manage members, attendance logs, and day-pass transactions, and share reports/analytics with Admin. "
        "Admin additionally views the occupancy dashboard, which depends on generated occupancy predictions fed by stored attendance data.",
        align="justify",
    )
    add_figure(doc, ASSETS / "01_usecase.png", "Fig. 1. GYMTRACE use case diagram.", width=Inches(3.35))

    add_para(doc, "Use Case UC-01: Authenticate via Facial Recognition (Member)", size=10, bold=True, first_line=False, space_after=3)
    add_usecase_table(
        doc,
        [
            ("Use Case ID / Name", "UC-01 Authenticate via Facial Recognition"),
            ("Actor", "Member"),
            ("Precondition", "Member is enrolled with an active membership and face profile."),
            ("Main flow", "Start check-in → capture face → match gallery → verify membership → log entry → update occupancy → show access granted."),
            ("Alternate flow", "No face / unknown face / inactive membership → deny access with reason."),
            ("Postcondition", "Attendance stored; occupancy forecast refreshed on success."),
        ],
    )

    add_para(doc, "Use Case UC-02: Purchase / Enter as Day-Pass User", size=10, bold=True, first_line=False, space_after=3)
    add_usecase_table(
        doc,
        [
            ("Use Case ID / Name", "UC-02 Purchase Day Pass / Enter Gym as Day-Pass User"),
            ("Actor", "Day-Pass User"),
            ("Precondition", "Visitor provides identity/contact and completes mock payment, or presents a valid unused pass code."),
            ("Main flow", "Purchase pass → present/validate pass → log attendance → allow entry."),
            ("Alternate flow", "Invalid or already-used pass → deny access."),
            ("Postcondition", "Day-pass marked used; attendance stored."),
        ],
    )

    add_para(doc, "Use Case UC-03: Staff Operations & UC-04 Admin Analytics", size=10, bold=True, first_line=False, space_after=3)
    add_usecase_table(
        doc,
        [
            ("UC-03 Staff", "Manage members (enroll/update), view attendance logs, manage day-pass transactions, view reports."),
            ("UC-04 Admin", "View reports & analytics and occupancy dashboard driven by Generate Occupancy Prediction."),
            ("Includes", "Log Attendance → Store Attendance Data → Generate Occupancy Prediction."),
        ],
    )

    # B. Activity
    add_heading_ieee(doc, "B. Activity Diagram per Use Case", level=2)
    add_para(
        doc,
        "Fig. 2 presents the activity diagram for the combined Gym Entry & Occupancy Update flow (covering UC-01 and UC-02). "
        "After selecting entry type, day-pass users validate a pass while members capture and match a face, then verify membership status. "
        "Both success paths converge on log attendance, store data, generate occupancy prediction, and allow access; failure paths deny access with a reason.",
        align="justify",
    )
    add_figure(doc, ASSETS / "02_activity.png", "Fig. 2. Activity diagram: gym entry and occupancy update.", width=Inches(3.35))

    # C. Class
    add_heading_ieee(doc, "C. Class Diagram", level=2)
    add_para(
        doc,
        "Fig. 3 shows the structural design. User specializes into Staff and Admin. Member associates 1:1 with Membership and FaceProfile and 1:* with AttendanceRecord. "
        "DayPassUser purchases DayPass objects that may create an AttendanceRecord. "
        "AuthenticationService uses Membership and FaceProfile. "
        "OccupancyService manages OccupancyForecast objects fed by attendance records. "
        "This design separates identity, access credentials, attendance events, and forecasting concerns.",
        align="justify",
    )
    add_figure(doc, ASSETS / "03_class.png", "Fig. 3. GYMTRACE class diagram.", width=Inches(3.35))

    # D. Sequence
    add_heading_ieee(doc, "D. Sequence Diagram per Use Case", level=2)
    add_para(
        doc,
        "Fig. 4 details the main success scenario of UC-01 (Authenticate via Facial Recognition). "
        "The Member interacts with CheckInUI; AuthService captures and authenticates the face through FaceProfile matching, verifies Membership status, logs AttendanceRecord, requests OccupancyService.generatePrediction(), and returns accessGranted to the UI. "
        "Alternate flows return accessDenied(reason) for unknown faces or inactive memberships.",
        align="justify",
    )
    add_figure(doc, ASSETS / "04_sequence.png", "Fig. 4. Sequence diagram: authenticate via facial recognition.", width=Inches(3.35))

    # E. Architecture
    add_heading_ieee(doc, "E. System Architecture (3-Tier)", level=2)
    add_para(
        doc,
        "Fig. 5 illustrates the GYMTRACE three-tier architecture with edge devices and an ML layer. "
        "Presentation tier: Admin UI, Member UI, Day-Pass UI, and Staff UI. "
        "Application tier: API Gateway routing to Authentication, Member & Day-Pass, Attendance, Occupancy, and Reporting services. "
        "Data tier: Users & Memberships DB, Face Embeddings Store, Attendance Logs DB, and Occupancy Forecasts DB. "
        "Edge cameras/kiosks feed face frames to Authentication; ML components provide facial recognition and occupancy prediction models consumed by application services.",
        align="justify",
    )
    add_figure(doc, ASSETS / "05_architecture.png", "Fig. 5. GYMTRACE three-tier system architecture.", width=Inches(3.35))

    # ===================== SECTION III =====================
    add_heading_ieee(doc, "III. USER INTERFACE DESIGN")
    add_para(
        doc,
        "The GYMTRACE prototype GUI is a dark athletic-themed Next.js web application that mirrors the actors and use cases in Section II. "
        "Fig. 6–7 show early wireframes and information structure. Figs. 8–23 are screenshots of the actual running GUI covering every primary screen: home, member flows, day-pass flows, staff console, and admin analytics.",
        align="justify",
    )
    add_figure(doc, ASSETS / "06_ui_wireframes.png", "Fig. 6. Frontend wireframes for role-based GYMTRACE screens.", width=Inches(3.35))
    add_figure(doc, ASSETS / "07_ui_structure.png", "Fig. 7. Frontend information structure by role.", width=Inches(3.2))

    gui = ASSETS / "gui"
    screens = [
        (
            "gui_01_home.jpg",
            "Fig. 8. Actual GUI: home role-select screen.",
            "Home. The landing page presents four access lanes—Member, Day-Pass, Staff, and Admin—so kiosk users immediately choose the correct workflow.",
        ),
        (
            "gui_02_member_checkin.jpg",
            "Fig. 9. Actual GUI: member face check-in.",
            "Member — Face Check-In. Camera viewport (CAM-01), scanner status, Start Face Scan, plus shortcuts to occupancy and attendance history.",
        ),
        (
            "gui_05_member_result.jpg",
            "Fig. 10. Actual GUI: member access-granted result.",
            "Member — Result. After a successful scan the UI shows CLEAR / Access Granted with the matched member name and links to occupancy and history.",
        ),
        (
            "gui_03_member_occupancy.jpg",
            "Fig. 11. Actual GUI: member predicted occupancy.",
            "Member — Predicted Occupancy. Live headcount, next-30-minute forecast, capacity load, and a forecast chart from the occupancy service.",
        ),
        (
            "gui_04_member_history.jpg",
            "Fig. 12. Actual GUI: member attendance history.",
            "Member — Attendance History. Personal visit list with timestamps and FACE entry badges.",
        ),
        (
            "gui_06_daypass_purchase.jpg",
            "Fig. 13. Actual GUI: purchase day pass.",
            "Day-Pass — Purchase. Name/contact form, GCash/Card mock payment, ₱250 price, and Pay & Issue Pass.",
        ),
        (
            "gui_08_daypass_issued.jpg",
            "Fig. 14. Actual GUI: day pass issued.",
            "Day-Pass — Issued. Displays the access code for the entrance kiosk and Continue to Entry.",
        ),
        (
            "gui_07_daypass_entry.jpg",
            "Fig. 15. Actual GUI: day-pass entry validation.",
            "Day-Pass — Entry. Pass-code field and Validate & Enter to log guest attendance.",
        ),
        (
            "gui_09_daypass_result.jpg",
            "Fig. 16. Actual GUI: day-pass entry allowed.",
            "Day-Pass — Result. Success state shows GO / Entry Allowed with pass code and guest name.",
        ),
        (
            "gui_10_staff_console.jpg",
            "Fig. 17. Actual GUI: staff console home.",
            "Staff — Console. PIN-unlocked sidebar (Console, Members, Attendance, Day Passes) with tiles for enrollment, logs, and guest-pass ops.",
        ),
        (
            "gui_11_staff_members.jpg",
            "Fig. 18. Actual GUI: staff members management.",
            "Staff — Members. Enroll form plus member table with status badges (active/expired/frozen), face tags, and status actions.",
        ),
        (
            "gui_12_staff_attendance.jpg",
            "Fig. 19. Actual GUI: staff attendance logs.",
            "Staff — Attendance. Recent face and day-pass entries with timestamps and source badges.",
        ),
        (
            "gui_13_staff_daypasses.jpg",
            "Fig. 20. Actual GUI: staff day-pass sales.",
            "Staff — Day Passes. Sales table with code, guest, amount, paid time, and used/available status.",
        ),
        (
            "gui_14_admin_home.jpg",
            "Fig. 21. Actual GUI: admin home.",
            "Admin — Home. Workspace picker separating Occupancy & ML from Revenue so analytics stay clear.",
        ),
        (
            "gui_15_admin_occupancy.jpg",
            "Fig. 22. Actual GUI: admin occupancy & ML dashboard.",
            "Admin — Occupancy. Live load, model metrics (MAE/RMSE/R²/MAPE), 12-hour forecast, actual-vs-predicted, peak hours, and feature importance charts.",
        ),
        (
            "gui_16_admin_revenue.jpg",
            "Fig. 23. Actual GUI: admin revenue dashboard.",
            "Admin — Revenue. Membership and day-pass income KPIs, monthly totals, and 7-day revenue/entries charts kept separate from occupancy views.",
        ),
    ]

    for fname, caption, explain in screens:
        add_para(doc, explain, align="justify")
        add_figure(doc, gui / fname, caption, width=Inches(3.35))

    add_para(
        doc,
        "Overall, the UI design prioritizes fast kiosk interaction, clear success/failure feedback for authentication, and strict role separation so members never see staff or admin controls.",
        align="justify",
    )

    # Acknowledgment + References
    add_heading_ieee(doc, "ACKNOWLEDGMENT")
    add_para(
        doc,
        "The authors thank Mapua Malayan Colleges Mindanao and the College of Computer and Information Science for academic support. Any opinions and conclusions are those of the authors.",
        align="justify",
    )

    add_heading_ieee(doc, "REFERENCES")
    refs = [
        'Q. Meng, S. Zhao, Z. Huang, and F. Zhou, “MagFace: A universal representation for face recognition and quality assessment,” in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2021, pp. 14220–14229, doi: 10.1109/CVPR46437.2021.01400.',
        'M. Kim, A. K. Jain, and X. Liu, “AdaFace: Quality adaptive margin for face recognition,” in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2022, pp. 18729–18738, doi: 10.1109/CVPR52688.2022.01819.',
        'F. Boutros, N. Damer, F. Kirchbuchner, and A. Kuijper, “ElasticFace: Elastic margin loss for deep face recognition,” in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. Workshops (CVPRW), 2022, pp. 1577–1586, doi: 10.1109/CVPRW56347.2022.00164.',
        'A. Chowanda, J. Moniaga, J. C. Bahagiono, and J. S. Chandra, “Machine learning face recognition model for employee tracking and attendance system,” in Proc. Int. Conf. Inf. Manage. Technol. (ICIMTech), 2022, pp. 297–301, doi: 10.1109/ICIMTech55957.2022.9915078.',
        'K. Kushwaha, S. Rahul, S. Eliyaz, C. Reddy, K. Amarendra, and T. K. Rama Krishna Rao, “A CNN based attendance management system using face recognition,” in Proc. 4th Int. Conf. Smart Electron. Commun. (ICOSEC), 2023, pp. 880–884, doi: 10.1109/ICOSEC58147.2023.10276353.',
        'M. S. Ejaz, S. Debnath, M. K. Hasan, and M. M. Alam, “Facial recognition-based entry system for student residence halls: Enhancing security and accessibility,” Asian J. Res. Comput. Sci., vol. 16, no. 4, pp. 344–353, 2023, doi: 10.9734/ajrcos/2023/v16i4396.',
        'Z. D. Tekler and A. Chong, “Occupancy prediction using deep learning approaches across multiple space types: A minimum sensing strategy,” Building and Environment, vol. 226, Art. no. 109689, Dec. 2022, doi: 10.1016/j.buildenv.2022.109689.',
        'E. Hitimana, G. Bajpai, R. Musabe, L. Sibomana, and J. Kayalvizhi, “Implementation of IoT framework with data analysis using deep learning methods for occupancy prediction in a building,” Future Internet, vol. 13, no. 3, Art. no. 67, 2021, doi: 10.3390/fi13030067.',
        'C. Kanthila, A. Boodi, K. Beddiar, Y. Amirat, and M. Benbouzid, “Occupancy prediction in buildings using cascaded LSTM model,” in Proc. IECON 2023—49th Annu. Conf. IEEE Ind. Electron. Soc., 2023, pp. 1–6, doi: 10.1109/IECON51785.2023.10311629.',
        'M. K. Diarra, A. Maniar, J.-B. Masson, B. Marhic, and L. Delahoche, “Occupancy state prediction by recurrent neural network (LSTM): Multi-room context,” Sensors, vol. 23, no. 23, Art. no. 9603, 2023, doi: 10.3390/s23239603.',
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(f"[{i}] {ref}")
        set_run_font(run, size=8)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("Wrote", OUT)
    print("Size MB:", round(OUT.stat().st_size / 1e6, 2))


if __name__ == "__main__":
    build()
